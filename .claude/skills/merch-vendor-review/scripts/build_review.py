#!/usr/bin/env python3
"""
Build the ONE output of a review: a single workbook that answers "can we integrate?" and
"does the API work?".

Usage:
    python3 build_review.py <test_cases.json> <review.json> <Vendor>_Review.xlsx [--md <report.md>]

Prints a SHORT SUMMARY first (context, what we tested, what their API has, what is missing and why,
what needs to be done; ready to paste into Jira), then TWO boxed reports for the chat reply: the review (decision, must-haves, areas, vendor
questions) and the build feasibility in merch-sensei's format (AUTH, ORDER, IDEMPOTENCY, ...,
build files, handoff).

Workbook tabs:
    Summary           the answer (can we integrate? does the API work?), the 7 must-haves,
                      blockers, next steps
    Ask the vendor    numbered questions (must-haves first) and the email draft
    Tech              build feasibility, endpoints, field and status mapping, what can go wrong,
                      assumptions, setup changes (only when review.json has "tech")
    All checks        every check: result and what we found (must-haves first, filterable)

<review.json>:
    {
      "vendor_name": "Example Vendor",
      "reviewed_on": "2026-10-18",
      "doc_sources": ["https://vendor.example/docs"],
      "answers": {
        "ORD-02": {"provided": "Yes", "details": "Response field orders_id",
                   "live": "Pass", "live_notes": "200, orders_id present"},
        "WHK-05": {"provided": "No", "details": "Signature header is mandatory"}
      },
      "status_map": {"in production": "CONFIRMED", "dispatched": "SHIPPED"},
      "security": "No credentials found in the document.",                 # optional
      "blockers": ["Every error is HTTP 500 ...", "..."],                  # optional, short plain lines
      "api_note": "Order, status and catalogue worked; errors return HTTP 500.",  # optional, plain words
      "short_summary": {                                                   # optional; printed first
        "context": "Manual vendor today; we read the API doc and tested staging.",
        "tested": ["Getting the product list ✅", "Placing a test order (preview) ✅"],
        "has": [{"need": "Place order", "api": "?order", "status": "✅ Works"}],
        "missing": [{"title": "Orders are put on hold", "detail": "...", "why": "...", "blocker": true}],
        "todo": [{"task": "Approve our orders automatically", "who": "Vendor"}]
      },
      "must_have_notes": {"1": "orders_id returned", "4": "Only 1 status documented"},  # optional, by must-have #
      "questions": [{"q": "Full list of order statuses?", "need": "Must-have", "checks": "STS-01"}],  # optional
      "vendor_email": "Subject: ...\n\nHi ...",                              # optional
      "next_steps": [{"step": "Send questions", "who": "Us"}],             # optional
      "tech": {                                                            # optional
        "api_summary": "...",
        "endpoints": [{"operation": "", "call": "", "notes": ""}],
        "order_mapping": [{"their_field": "", "our_source": "", "rule": ""}],
        "status_flow": [{"their_status": "", "our_status": "", "effect": ""}],
        "failure_modes": [{"failure": "", "handling": "", "severity": ""}],
        "assumptions": [{"assumption": "", "verify": ""}],
        "config": [{"where": "", "change": ""}],
        "build_notes": ["..."],
        "class_name": "Printstop", "model": "printstop_merchandise", "build_status": "NEW or DIFF",
        "complexity": "Medium (why)", "risks": ["..."],
        "feasibility": {"AUTH": {"status": "⚠️", "note": "..."}}   # optional overrides per line
      }
    }

    provided: Yes / No / Not available / NA / Open      (missing = "Not available")
      Yes            docs confirm it; details MUST say where (endpoint, field, short paraphrase)
      No             docs confirm it is absent; details MUST say where
      Not available  docs are silent (the default; never guessed)
      NA             does not apply to this vendor
      Open           partly answered; needs the vendor to confirm
    live: Pass / Fail / Not run / NA    (sandbox result, only for checks marked Live)
    status_map: vendor status -> PENDING / CONFIRMED / SHIPPED / DELIVERED / CANCELED / RETURNED

Decision (from the Must-have checks only; Important and Nice-to-have never block):
    No             a must-have is missing, confirmed absent, or failed on sandbox
    Likely yes     no must-have is missing, but some are only partly answered (vendor to confirm)
    Yes on paper   every must-have is confirmed in the docs; sandbox test still to run
    Yes            every must-have is confirmed and passed on sandbox

Exits 2 on invalid input (unknown id, bad value, Yes/No without evidence, bad status target).
"""
import json
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY, GREY, WHITE = "1F3864", "F2F2F2", "FFFFFF"
GREEN, AMBER, RED = "C6E0B4", "FFE699", "F8CBAD"
PROVIDED = ("Yes", "No", "Not available", "NA", "Open")
LIVE = ("Pass", "Fail", "Not run", "NA")
DELIVERY_STATUSES = ("PENDING", "CONFIRMED", "SHIPPED", "DELIVERED", "CANCELED", "RETURNED")
NEEDS = ("Must-have", "Important", "Nice-to-have")
FILL = {"Yes": GREEN, "Pass": GREEN, "Met": GREEN, "No": RED, "Fail": RED, "Missing": RED,
        "Open": AMBER, "Not available": AMBER, "Not run": AMBER}


def font(**kw):
    return Font(name="Arial", **kw)


def border():
    s = Side(style="thin", color="D9D9D9")
    return Border(left=s, right=s, top=s, bottom=s)


def has_live(t):
    return "Live" in t["verify"]


# ---------- validation ----------

def validate(tests, payload):
    errors, warnings = [], []
    by_id = {t["id"]: t for t in tests}
    for tid, a in payload.get("answers", {}).items():
        if tid not in by_id:
            errors.append(f"{tid}: unknown check id")
            continue
        p = a.get("provided", "Not available")
        if p not in PROVIDED:
            errors.append(f"{tid}: provided '{p}' not in {PROVIDED}")
        if p in ("Yes", "No") and not a.get("details", "").strip():
            errors.append(f"{tid}: '{p}' needs evidence in details (never assume)")
        if p == "Open" and not a.get("details", "").strip():
            warnings.append(f"{tid}: Open without a note on what is unclear")
        live = a.get("live")
        if live is not None:
            if live not in LIVE:
                errors.append(f"{tid}: live '{live}' not in {LIVE}")
            elif not has_live(by_id[tid]) and live != "NA":
                warnings.append(f"{tid}: sandbox result given for a docs-only check")
    for status, target in payload.get("status_map", {}).items():
        if target not in DELIVERY_STATUSES:
            errors.append(f"status_map '{status}' -> '{target}': must be one of {DELIVERY_STATUSES}"
                          " (failed/rejected map to CANCELED)")
    return errors, warnings


# ---------- evaluation ----------

def plural(n, word):
    return f"{n} {word}{'' if n == 1 else 's'}"


def evaluate(data, answers, status_map):
    tests = data["tests"]
    by_id = {t["id"]: t for t in tests}

    def provided(tid):
        a = answers.get(tid, {})
        if "provided" not in a and by_id[tid]["verify"] == "Live":
            # sandbox-only checks are answered by the sandbox run
            return {"Pass": "Yes", "Fail": "No", "NA": "NA"}.get(a.get("live"), "Not available")
        return a.get("provided", "Not available")

    def live(tid):
        if not has_live(by_id[tid]) or provided(tid) == "NA":
            return "NA"
        return answers.get(tid, {}).get("live", "Not run")

    def met(tid):
        return provided(tid) == "Yes" and live(tid) != "Fail"

    def state(tid):  # Met / Open / Missing
        if met(tid) or provided(tid) == "NA":
            return "Met"
        if provided(tid) == "Open" and live(tid) != "Fail":
            return "Open"
        return "Missing"

    must = []
    for g in data["must_haves"]:
        all_states = {tid: state(tid) for tid in g.get("all", [])}
        any_ids = g.get("any", [])
        any_states = {tid: state(tid) for tid in any_ids}
        failed = [tid for tid in list(all_states) + any_ids if live(tid) == "Fail"]
        if failed or "Missing" in all_states.values() or (any_ids and set(any_states.values()) == {"Missing"}):
            s = "Missing"  # a sandbox failure on any must-have check (even an alternative) blocks
        elif "Open" in all_states.values() or (any_ids and "Met" not in any_states.values()):
            s = "Open"
        else:
            s = "Met"
        ids = list(all_states) + any_ids
        evidence = "\n".join(f"{tid}: {answers.get(tid, {}).get('details', '') or provided(tid)}"
                              for tid in ids if provided(tid) != "Not available" or tid in all_states)
        must.append({"item": g["item"], "short": g.get("short", g["item"]), "state": s, "ids": ids,
                     "evidence": evidence})

    counts = {s: sum(m["state"] == s for m in must) for s in ("Met", "Open", "Missing")}

    live_needed = [t for t in tests if live(t["id"]) != "NA"]
    live_pass = [t for t in live_needed if live(t["id"]) == "Pass"]
    live_fail = [t for t in live_needed if live(t["id"]) == "Fail"]
    must_ids = {i for m in must for i in m["ids"]}
    must_live_open = [t for t in live_needed if t["id"] in must_ids and met(t["id"]) and live(t["id"]) != "Pass"]
    if not live_pass and not live_fail:
        api_works = "Not tested yet (no sandbox run)"
    elif live_fail:
        api_works = f"Problems found: {len(live_fail)} failed on sandbox ({', '.join(t['id'] for t in live_fail)})"
    elif len(live_pass) == len(live_needed):
        api_works = f"Works on sandbox: all {len(live_pass)} sandbox checks passed"
    else:
        api_works = f"Partly tested: {len(live_pass)} of {len(live_needed)} sandbox checks passed, none failed"

    escalations = [t for t in tests if t.get("risk") and provided(t["id"]) in ("No", "Open")]

    if counts["Missing"]:
        decision = "No, we cannot integrate yet"
        reason = f"{plural(counts['Missing'], 'must-have')} missing"
    elif counts["Open"]:
        decision = "Likely yes, after the vendor answers a few points"
        reason = f"{plural(counts['Open'], 'must-have')} only partly documented"
    elif must_live_open:
        decision = "Yes on paper, confirm with a sandbox test"
        reason = "every must-have is documented; sandbox test still to run"
    else:
        decision = "Yes, we can integrate"
        reason = "every must-have is documented and passed on sandbox"
    if escalations and not counts["Missing"]:
        n = len(escalations)
        reason += f"; {plural(n, 'item')} {'needs' if n == 1 else 'need'} a change on our side first"

    def gaps(need):
        return [t for t in tests if t["need"] == need and state(t["id"]) != "Met"]

    return {
        "provided": provided, "live": live, "state": state, "must": must, "counts": counts,
        "decision": decision, "reason": reason, "api_works": api_works, "live_fail": live_fail,
        "escalations": escalations, "gaps": {n: gaps(n) for n in NEEDS},
        "covered": {n: (sum(state(t["id"]) == "Met" for t in tests if t["need"] == n),
                        sum(t["need"] == n for t in tests)) for n in NEEDS},
    }


# ---------- workbook (clean layout: one question per tab) ----------

TECH_SECTIONS = [  # key, title, [(field, column label)]
    ("endpoints", "Endpoints", [("operation", "Operation"), ("call", "Call"), ("notes", "Notes")]),
    ("order_mapping", "Order fields: ours to theirs", [("our_source", "Our field"), ("their_field", "Their field"),
                                                       ("rule", "Rule")]),
    ("status_flow", "Status mapping", [("their_status", "Their status"), ("our_status", "Our status"),
                                       ("effect", "What happens")]),
    ("failure_modes", "What can go wrong", [("failure", "Problem"), ("severity", "Severity"),
                                            ("handling", "How we handle it")]),
    ("assumptions", "Assumptions to check", [("assumption", "Assumption"), ("verify", "How to check")]),
    ("config", "Setup changes", [("where", "Where"), ("change", "Change")]),
]
STATE_LABEL = {"Met": "✓  Yes", "Open": "!  Partly", "Missing": "✗  No"}
ANSWER_LABEL = {"Yes": "✓  Yes", "Open": "!  Partly", "No": "✗  No", "Not available": "–  Not stated", "NA": "n/a"}
ANSWER_FILL = {"Yes": GREEN, "Open": AMBER, "No": RED, "Not available": "EDEDED"}
LIGHT = "D9E1F2"


class Sheet:
    """Writes clean blocks top to bottom: title, section headings, tables, label/value rows, text boxes."""

    def __init__(self, ws, widths):
        self.ws, self.row, self.ncols = ws, 1, len(widths)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 110

    def _cell(self, r, c, value, **style):
        cell = self.ws.cell(row=r, column=c, value=value)
        cell.font = font(**{k: v for k, v in style.items() if k in ("bold", "size", "color", "italic")})
        cell.alignment = Alignment(wrap_text=True, vertical="center" if style.get("center") else "top",
                                   horizontal="center" if style.get("center") else "left")
        if style.get("fill"):
            cell.fill = PatternFill("solid", fgColor=style["fill"])
        if style.get("box"):
            cell.border = border()
        return cell

    def _span(self):
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncols)

    def title(self, text, sub=None):
        self._span()
        self._cell(self.row, 1, text, bold=True, size=18, color=NAVY)
        self.ws.row_dimensions[self.row].height = 30
        self.row += 1
        if sub:
            self._span()
            self._cell(self.row, 1, sub, italic=True, color="595959")
            self.row += 1

    def heading(self, text):
        self.row += 1
        self._span()
        self._cell(self.row, 1, text, bold=True, size=12, color=WHITE, fill=NAVY)
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def table(self, labels, rows, first_col_fills=None, center_first=False):
        for c, label in enumerate(labels, 1):
            self._cell(self.row, c, label, bold=True, fill=LIGHT, box=True)
        self.row += 1
        for i, values in enumerate(rows):
            for c, v in enumerate(values, 1):
                fill = first_col_fills[i] if (first_col_fills and c == 1) else None
                self._cell(self.row, c, v, box=True, fill=fill, bold=(c == 1 and bool(first_col_fills)),
                           center=(c == 1 and center_first))
            self.row += 1

    def big(self, label, value, fill, note=None):
        """A large answer row: LABEL | BIG VALUE (merged across), with an optional note row below."""
        self._cell(self.row, 1, label, bold=True, size=12, center=True)
        self.ws.merge_cells(start_row=self.row, start_column=2, end_row=self.row, end_column=self.ncols)
        self._cell(self.row, 2, value, bold=True, size=16, fill=fill, center=True)
        self.ws.row_dimensions[self.row].height = 34
        self.row += 1
        if note:
            self.ws.merge_cells(start_row=self.row, start_column=2, end_row=self.row, end_column=self.ncols)
            self._cell(self.row, 2, note, italic=True, color="404040")
            self.row += 1

    def text(self, value, height=None):
        self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.ncols)
        self._cell(self.row, 1, value, box=True)
        if height:
            self.ws.row_dimensions[self.row].height = height
        self.row += 1

    def bullets(self, items):
        for i, item in enumerate(items, 1):
            self._cell(self.row, 1, i, bold=True, center=True)
            self.ws.merge_cells(start_row=self.row, start_column=2, end_row=self.row, end_column=self.ncols)
            self._cell(self.row, 2, item)
            self.row += 1


def vendor_questions(data, payload, ev):
    """Questions from review.json, else generated from open/missing checks (must-haves first)."""
    if payload.get("questions"):
        return [(q.get("need", ""), q["q"], q.get("checks", "")) for q in payload["questions"]]
    rows = []
    for m in ev["must"]:
        if m["state"] != "Met":
            ids = [i for i in m["ids"] if ev["state"](i) != "Met"]
            rows.append(("Must-have", m["item"], ", ".join(ids)))
    for t in ev["gaps"]["Important"]:
        rows.append(("Important", t["test"], t["id"]))
    return rows


def headline(ev):
    d = ev["decision"]
    if d.startswith("No"):
        return "NO, NOT YET", RED
    if d.startswith("Likely"):
        return "LIKELY YES", AMBER
    if d.startswith("Yes on paper"):
        return "YES, AFTER A SANDBOX TEST", AMBER
    return "YES", GREEN


def api_headline(ev):
    a = ev["api_works"]
    if a.startswith("Works"):
        return "YES, ALL SANDBOX TESTS PASSED", GREEN
    if a.startswith("Problems"):
        return f"MOSTLY: {len(ev['live_fail'])} SANDBOX ISSUE{'S' if len(ev['live_fail']) != 1 else ''}", AMBER
    if a.startswith("Partly"):
        return "PARTLY TESTED", AMBER
    return "NOT TESTED YET", "EDEDED"


def api_note(data, ev):
    run = [t for t in data["tests"] if ev["live"](t["id"]) in ("Pass", "Fail")]
    if not run:
        return "No sandbox tests yet. Ask the vendor for sandbox access, then run LIVE mode."
    fails = len(ev["live_fail"])
    note = f"{len(run) - fails} of {len(run)} sandbox checks passed"
    return note + (f"; {fails} failed (marked red on 'All checks')." if fails else ".")


def build_summary(wb, data, payload, ev):
    ws = wb.active
    ws.title = "Summary"
    s = Sheet(ws, [24, 34, 74])
    s.title(payload.get("vendor_name", "Vendor"), f"Merchandise API review · {payload.get('reviewed_on', '')}")

    s.heading("The answer")
    text, fill = headline(ev)
    s.big("Can we integrate?", text, fill, ev["reason"][:1].upper() + ev["reason"][1:] + ".")
    text, fill = api_headline(ev)
    s.big("Does the API work?", text, fill, payload.get("api_note") or api_note(data, ev))

    c = ev["counts"]
    s.heading(f"Must-haves: {c['Met']} of {len(ev['must'])} met")
    notes = payload.get("must_have_notes", {})
    rows = [(STATE_LABEL[m["state"]], m["item"], notes.get(str(i), m["evidence"].split("\n")[0]))
            for i, m in enumerate(ev["must"], 1)]
    s.table(["Result", "Must-have", "What we found"], rows,
            first_col_fills=[FILL[m["state"]] for m in ev["must"]], center_first=True)

    blockers = payload.get("blockers") or (payload.get("tech") or {}).get("risks") or []
    if payload.get("security") and not payload["security"].lower().startswith("no credentials"):
        blockers = [f"Security: {payload['security']}"] + list(blockers)
    if blockers:
        s.heading("Blockers to clear before go-live")
        s.bullets(blockers)

    steps = payload.get("next_steps") or []
    if steps:
        s.heading("Next steps")
        s.bullets([f"{x['step']}  ({x['who']})" if x.get("who") else x["step"] for x in steps])

    s.row += 1
    s._cell(s.row, 1, "More detail: 'Ask the vendor' (questions + email) · 'Tech' (build view) · "
                      "'All checks' (every check with evidence)", italic=True, color="7F7F7F")
    ws.merge_cells(start_row=s.row, start_column=1, end_row=s.row, end_column=3)


def build_questions(wb, data, payload, ev):
    ws = wb.create_sheet("Ask the vendor")
    s = Sheet(ws, [6, 14, 112])
    s.title("Ask the vendor", "Must-haves first. Other questions are good to know but do not block.")
    qs = vendor_questions(data, payload, ev)
    s.heading("Questions")
    fills = [RED if q[0] == "Must-have" else None for q in qs]
    s.table(["#", "Priority", "Question"], [(i, q[0], q[1]) for i, q in enumerate(qs, 1)], center_first=True)
    for r, f in zip(range(s.row - len(qs), s.row), fills):
        if f:
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=AMBER)
    if payload.get("vendor_email"):
        s.heading("Email draft (ready to send)")
        s.text(payload["vendor_email"], height=max(160, 16 * (payload["vendor_email"].count("\n") + 8)))


def build_tech(wb, payload, ev):
    tech = payload.get("tech")
    if not tech:
        return
    ws = wb.create_sheet("Tech")
    s = Sheet(ws, [30, 40, 80])
    s.title("Tech", "For developers: how this vendor fits our merchandise platform.")
    if tech.get("api_summary"):
        s.heading("How the API works")
        s.text(tech["api_summary"], height=75)
    s.heading("Build feasibility")
    word = {"✅": ("✓  Ready", GREEN), "⚠️": ("!  Check", AMBER), "❌": ("✗  Blocked", RED), "➖": ("–  Not needed", "EDEDED")}
    rows = feasibility(payload, ev)
    s.table(["Area", "Status", "Note"], [(l.title().replace("/", " / "), word.get(i, (i, WHITE))[0], n) for l, i, n in rows])
    for r, (_, icon, _) in zip(range(s.row - len(rows), s.row), rows):
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=word.get(icon, (icon, WHITE))[1])
        ws.cell(row=r, column=2).font = font(bold=True)
    sm = payload.get("status_map", {})
    for key, title, cols in TECH_SECTIONS:
        rows = tech.get(key) or []
        if key == "status_flow" and not rows and sm:
            rows = [{"their_status": k, "our_status": v, "effect": ""} for k, v in sm.items()]
        if not rows:
            continue
        s.heading(title)
        labels = [label for _, label in cols]
        s.table(labels + [""] * (3 - len(labels)), [[r.get(f, "") for f, _ in cols] for r in rows])
    if tech.get("build_notes"):
        s.heading("Build notes")
        s.bullets(tech["build_notes"])


def build_checks(wb, data, answers, ev):
    ws = wb.create_sheet("All checks")
    cols = [("Area", 24), ("Need", 13), ("Check", 60), ("Result", 14), ("What we found", 70), ("ID", 9)]
    for c, (label, w) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.font = font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    order = {"Must-have": 0, "Important": 1, "Nice-to-have": 2}
    tests = sorted(data["tests"], key=lambda t: order[t["need"]])  # must-haves first, then the rest
    for t in tests:
        a = answers.get(t["id"], {})
        p, lv = ev["provided"](t["id"]), ev["live"](t["id"])
        found = a.get("details", "")
        if lv in ("Pass", "Fail"):
            found = (found + "\n" if found else "") + f"Sandbox {lv.lower()}: {a.get('live_notes', '')}".rstrip(": ")
        ws.append([t["section"], t["need"], t["test"], ANSWER_LABEL.get(p, p), found, t["id"]])
        r = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = font(bold=(c == 2 and t["need"] == "Must-have"), color="7F7F7F" if c == 6 else "000000")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border()
        if p in ANSWER_FILL:
            ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=ANSWER_FILL[p])
        if lv == "Fail":
            ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=RED)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.sheet_view.showGridLines = False


# ---------- short summary (first thing in every reply; also the Jira comment) ----------

def build_short_summary(payload, ev):
    """Plain summary: context, what we tested, what their API has, what is missing and why, what to do.
    Uses review.json "short_summary" when given; otherwise builds a simpler version from the review."""
    ss = payload.get("short_summary") or {}
    name = payload.get("vendor_name", "Vendor").split(" (")[0]
    headline_text = {"No": "we cannot connect yet", "Likely yes": "we can connect, after a few answers from them",
                     "Yes on paper": "we can connect, after one sandbox test", "Yes,": "we can connect"}
    verdict = next((v for k, v in headline_text.items() if ev["decision"].startswith(k)), ev["decision"])
    out = [f"**{name} API review: {verdict}**", ""]

    out += ["**Context**", ss.get("context") or f"We reviewed {name}'s API documentation"
            + ("." if ev["api_works"].startswith("Not tested") else " and tested their staging system."), ""]

    tested = ss.get("tested") or []  # only what was really run; never derived
    if tested:
        out += ["**What we tested**"] + [f"- {t}" for t in tested] + [""]

    has = ss.get("has") or [{"need": m["item"], "api": "", "status": {"Met": "✅ Yes", "Open": "⚠️ Partly",
                                                                       "Missing": "❌ No"}[m["state"]]}
                            for m in ev["must"]]
    out += ["**What their API has**", "| Need | Their API | Status |", "|---|---|---|"]
    out += [f"| {h['need']} | {('`' + h['api'] + '`') if h.get('api') else ''} | {h['status']} |" for h in has]
    out.append("")

    missing = ss.get("missing")
    if missing is None:
        missing = [{"title": b, "detail": "", "why": "", "blocker": i == 0}
                   for i, b in enumerate(payload.get("blockers") or (payload.get("tech") or {}).get("risks") or [])]
    if missing:
        out.append("**What is missing, and why it matters**")
        for i, m in enumerate(missing, 1):
            flag = "🚫 " if m.get("blocker") else ""
            tail = " (BLOCKER)" if m.get("blocker") else ""
            out += ["", f"{flag}**{i}. {m['title']}{tail}**"]
            if m.get("detail"):
                out.append(m["detail"])
            if m.get("why"):
                out.append(f"**Why it matters:** {m['why']}")
            if m.get("ref"):
                out.append(f"_{m['ref']}_")
        out.append("")

    todo = ss.get("todo") or [{"task": x["step"], "who": x.get("who", "")} for x in payload.get("next_steps") or []]
    if todo:
        out += ["**What needs to be done**", "| # | Task | Who |", "|---|---|---|"]
        out += [f"| {i} | {t['task']} | {t.get('who', '')} |" for i, t in enumerate(todo, 1)]
    return "\n".join(out).rstrip() + "\n"


# ---------- terminal report (merch-sensei style) ----------

ICON = {"Met": "✅", "Open": "⚠️", "Missing": "❌"}
BAR = "━" * 64


def clip(text, n=88):
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1] + "…"


def wrapped(label, icon, text, width=100, w=15):
    """'LABEL           ✅ text' with long text wrapped under itself.
    icon None = list item (indent under the label text), icon "" = aligned label without an icon."""
    import textwrap
    if icon is None:
        head = f"{label} "
    elif icon:
        head = f"{label:<{w}} {icon} "
    else:
        head = f"{label:<{w}}    "
    pad = " " * (w + 4) if icon is not None else " " * len(head)  # icons are 2 columns wide on screen
    body = textwrap.wrap(" ".join(str(text).split()), width=width - len(pad)) or [""]
    return [head + body[0]] + [pad + line for line in body[1:]]


def build_terminal_report(data, payload, ev, out_path=None):
    """The boxed report the agent prints in chat, in the same style as merch-sensei."""
    c = ev["counts"]
    notes = payload.get("must_have_notes", {})
    verdict_icon = "❌" if c["Missing"] else "⚠️" if c["Open"] or ev["escalations"] else "✅"
    api_icon = "❌" if ev["live_fail"] else "✅" if ev["api_works"].startswith("Works") else "⚠️"
    lines = [BAR, f"🛒 Merch Vendor Review: {payload.get('vendor_name', 'Vendor')}", BAR,
             f"{'VERDICT':<15} {verdict_icon} {ev['decision']}",
             f"{'':<15}    {ev['reason']}",
             f"{'API TEST':<15} {api_icon} {ev['api_works']}"]
    if payload.get("security"):
        sec_icon = "❌" if payload["answers"].get("SEC-07", {}).get("provided") == "No" else "✅"
        lines += wrapped("SECURITY", sec_icon, payload["security"])
    lines += ["", f"MUST-HAVES ({c['Met']}/{len(ev['must'])} met · {c['Open']} open · {c['Missing']} missing)"]
    for i, m in enumerate(ev["must"], 1):
        note = notes.get(str(i)) or m["evidence"].split("\n")[0]
        lines += wrapped(m["short"], ICON[m["state"]], note)
    if ev["escalations"]:
        lines += ["", "NEEDS A CHANGE ON OUR SIDE"]
        for t in ev["escalations"]:
            lines += wrapped(t["id"], "⚠️", t["test"])
    lines += ["", "AREAS (checks answered Yes / applicable)"]
    tests = data["tests"]
    for section in dict.fromkeys(t["section"] for t in tests):
        ts = [t for t in tests if t["section"] == section and ev["provided"](t["id"]) != "NA"]
        yes = sum(ev["state"](t["id"]) == "Met" for t in ts)
        icon = "✅" if ts and yes == len(ts) else "❌" if yes == 0 else "⚠️"
        lines.append(f"{clip(section, 34):<34} {icon} {yes}/{len(ts)}")
    imp, nice = ev["covered"]["Important"], ev["covered"]["Nice-to-have"]
    lines.append(f"{'GAPS':<15} Important {imp[0]}/{imp[1]} covered · Nice-to-have {nice[0]}/{nice[1]} (never block)")
    qs = vendor_questions(data, payload, ev)
    if qs:
        lines += ["", "ASK THE VENDOR (top 5)"]
        for i, q in enumerate(qs[:5], 1):
            lines += wrapped(f" {i}. [{q[0]}]", None, q[1])
    steps = payload.get("next_steps") or []
    if steps:
        lines += ["", "NEXT STEPS"]
        for i, x in enumerate(steps, 1):
            lines += wrapped(f" {i}.", None, f"{x['step']} ({x.get('who', '')})")
    if out_path:
        lines += ["", f"{'OUTPUT':<15} {out_path}"]
    lines.append(BAR)
    return "\n".join(lines)


FEAS_MAP = [  # merch-sensei feasibility lines → checks that decide them
    ("AUTH", ["AUTH-01", "AUTH-02", "AUTH-03"]),
    ("ORDER", ["ORD-01", "ORD-02", "ORD-04"]),
    ("IDEMPOTENCY", ["IDM-01", "IDM-02", "IDM-03"]),
    ("ADDRESS", ["ORD-05", "ORD-06", "ORD-08"]),
    ("STATUS MAP", ["STS-01", "STS-02"]),
    ("WEBHOOK", ["WHK-01", "WHK-03", "WHK-04", "WHK-05"]),
    ("POLLING", ["POL-01", "POL-02"]),
    ("PRE-CONFIRM FAIL", ["ORD-13"]),
    ("TRACKING", ["TRK-01", "TRK-02"]),
    ("CANCEL/RETURNS", ["RTN-01", "RTN-03"]),
    ("CATALOG", ["CAT-01", "CAT-02", "CAT-03", "CAT-04"]),
    ("BALANCE", ["BAL-01"]),
]


def feasibility(payload, ev):
    """[(line, icon, note)]: tech.feasibility overrides, otherwise derived from the checks."""
    override = (payload.get("tech") or {}).get("feasibility", {})
    answers = payload.get("answers", {})
    rows = []
    for line, ids in FEAS_MAP:
        if line in override:
            rows.append((line, override[line].get("status", "⚠️"), override[line].get("note", "")))
            continue
        states = [ev["state"](i) for i in ids]
        documented = [i for i in ids if ev["provided"](i) not in ("Not available",)]
        if all(st == "Met" for st in states):
            icon = "✅"
        elif not documented:
            icon = "❌"
        else:
            icon = "⚠️"
        note = next((answers[i]["details"] for i in ids if answers.get(i, {}).get("details")), "Not documented")
        rows.append((line, icon, note))
    by = {r[0]: r for r in rows}
    # webhook OR polling is enough; balance is never required
    if "WEBHOOK" not in override and by["WEBHOOK"][1] == "❌" and by["POLLING"][1] == "✅":
        by["WEBHOOK"] = ("WEBHOOK", "➖", "None documented; polling is used instead")
    if "POLLING" not in override and by["POLLING"][1] == "❌" and by["WEBHOOK"][1] == "✅":
        by["POLLING"] = ("POLLING", "➖", "Not needed; webhooks are used")
    if "BALANCE" not in override and by["BALANCE"][1] == "❌":
        by["BALANCE"] = ("BALANCE", "➖", "None (not needed; no karma_reports module)")
    return [by[line] for line, _ in FEAS_MAP]


def build_feasibility_report(data, payload, ev):
    """Second box: build feasibility in merch-sensei's format, for developers."""
    tech = payload.get("tech") or {}
    cls = tech.get("class_name") or "".join(
        w[:1].upper() + w[1:] for w in re.split(r"[^A-Za-z0-9]+", payload.get("vendor_name", "Vendor").split("(")[0]) if w)
    lines = [BAR, f"🛠️ Build Feasibility (merch-sensei format): {cls}", BAR]
    for line, icon, note in feasibility(payload, ev):
        lines += wrapped(line, icon, note, w=17)
    live_total = sum(ev["live"](t["id"]) != "NA" for t in data["tests"])
    live_run = sum(ev["live"](t["id"]) in ("Pass", "Fail") for t in data["tests"])
    lines += wrapped("SANDBOX RUN", "", f"{live_run}/{live_total} sandbox checks run; "
                     f"{len(ev['live_fail'])} failed", w=17)
    if tech.get("complexity"):
        lines += wrapped("COMPLEXITY", "", tech["complexity"], w=17)
    risks = tech.get("risks") or [f["failure"] for f in tech.get("failure_modes", []) if f.get("severity") == "High"]
    if risks:
        lines += ["", "RISKS"]
        for r in risks:
            lines += wrapped(" •", None, r)
    model = tech.get("model", "{model}")
    status = tech.get("build_status", "NEW")
    lines += ["", "BUILD FILES (vendor-integrations, tracked)",
              f" • src/vendors/integratedVendors/{cls}.ts".ljust(64) + f"[{status}]",
              f" • src/vendors/vendorTemplates/{cls[:1].lower() + cls[1:]}Types.ts".ljust(64) + "[NEW, if needed]",
              " • src/configs/baseConfigs.json.sample".ljust(64) + "[UPDATE mapping + tracking ids]",
              f" • tests/vendors/{cls.lower()}.*.test.ts + fixtures".ljust(64) + "[NEW]",
              "LOCAL / DEPLOY CONFIG (gitignored, never in a PR)",
              f" • src/configs/vendorConfigs/{cls}.json, baseConfigs.json mappings (per env)",
              "CHECKLIST (outside VI)",
              " • OPMS integratedVendors (string id), schedulers (token, sync, tracking)",
              f" • DB: vendor, vendor_to_model, product_api_model_names, {model}_products",
              " • karma admin page, menu, routes, permissions (check what already exists)"]
    if tech.get("build_notes"):
        lines += ["", "NOTES"]
        for n in tech["build_notes"]:
            lines += wrapped(" •", None, n)
    ready = ev["counts"]["Missing"] == 0 and ev["counts"]["Open"] == 0
    lines += ["", *wrapped("HANDOFF", "", "/merch-sensei ONBOARD: " + ("ready" if ready else
                                                                    "after the open must-haves are answered"), w=17)]
    lines.append(BAR)
    return "\n".join(lines)


# ---------- markdown (optional) ----------

def build_markdown(data, answers, payload, ev):
    def cell(v):
        return str(v).replace("|", "\\|").replace("\n", " ")

    c = ev["counts"]
    out = [f"# {payload.get('vendor_name', 'Vendor')}: merchandise API review", ""]
    if payload.get("reviewed_on"):
        out.append(f"Reviewed on: {payload['reviewed_on']}  ")
    for src in payload.get("doc_sources", []):
        out.append(f"Source: {src}  ")
    out += ["", "## Decision", "", f"**Can we integrate?** {ev['decision']} ({ev['reason']}).  ",
            f"**Does the API work?** {ev['api_works']}.  ",
            f"**Must-haves:** {c['Met']} met, {c['Open']} open, {c['Missing']} missing (of {len(ev['must'])})."]
    if payload.get("security"):
        out.append(f"  \n**Security:** {payload['security']}")
    out += ["", "## Must-haves", "", "| Status | Must-have | What we found |", "|---|---|---|"]
    out += [f"| {m['state']} | {cell(m['item'])} | {cell(m['evidence'])} |" for m in ev["must"]]
    qs = vendor_questions(data, payload, ev)
    out += ["", "## Vendor questions", ""] + [f"{i}. [{q[0]}] {q[1]} ({q[2]})" for i, q in enumerate(qs, 1)]
    for need in ("Important", "Nice-to-have"):
        out += ["", f"## {need} gaps (do not block)", ""]
        out += [f"- **{t['id']}** {t['test']}" for t in ev["gaps"][need]] or ["None."]
    return "\n".join(out) + "\n"


def main():
    args = sys.argv[1:]
    md_path = None
    if "--md" in args:
        i = args.index("--md")
        md_path = args[i + 1]
        del args[i:i + 2]
    if len(args) < 3:
        print(__doc__)
        sys.exit(1)
    tests_path, review_path, out_path = args[:3]
    with open(tests_path) as f:
        data = json.load(f)
    with open(review_path) as f:
        payload = json.load(f)
    payload.setdefault("answers", {})
    errors, warnings = validate(data["tests"], payload)
    for w in warnings:
        print(f"WARN  {w}")
    if errors:
        for e in errors:
            print(f"ERROR {e}")
        sys.exit(2)
    answers = payload["answers"]
    ev = evaluate(data, answers, payload.get("status_map", {}))
    wb = Workbook()
    build_summary(wb, data, payload, ev)
    build_questions(wb, data, payload, ev)
    build_tech(wb, payload, ev)
    build_checks(wb, data, answers, ev)
    wb.save(out_path)
    if md_path:
        with open(md_path, "w") as f:
            f.write(build_markdown(data, answers, payload, ev))
        print(f"Wrote {md_path}")
    print(build_short_summary(payload, ev))
    print(build_terminal_report(data, payload, ev, out_path))
    print()
    print(build_feasibility_report(data, payload, ev))


if __name__ == "__main__":
    main()
